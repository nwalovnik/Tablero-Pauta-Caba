"""Normaliza los archivos de Pauta Publicitaria de BA Data (2017 en adelante) en una sola base.

Salidas:
  data/pauta.csv.gz   una fila por orden: anio, fecha, rubro, medio (normalizado), medio_fuente, importe
  pauta_data.json     agregados para el tablero (por año, rubro, mes y medio; nominal y real)

Reglas (ver README):
  - De cada recurso se usa el CSV; si no hay CSV (o el XLSX es el único con datos) se usa el XLSX.
    Los XLSX de 2021 traen día y mes invertidos y los CSV de 2021-2022 redondean a pesos enteros:
    por eso el orden día/mes y el formato de los importes se deducen archivo por archivo.
  - Cuando hay CSV y XLSX se controla que sumen lo mismo (diferencia < 0,1%).
  - Real: cada orden se deflacta con el IPCBA del mes, a pesos promedio del último año completo.
"""
import os, re, io, csv, gzip, json, unicodedata, datetime as dt, warnings
from collections import defaultdict, Counter
import openpyxl

warnings.filterwarnings('ignore')
HERE = os.path.dirname(os.path.abspath(__file__))
RAW = os.path.join(HERE, 'raw')
OUT_JSON = os.path.join(HERE, '..', 'pauta_data.json')
IPC_XLSX = os.path.join(HERE, 'IPCBA_empalmado.xlsx')

RUBROS = ['TV abierta', 'TV cable', 'Vía pública', 'Web', 'Radio', 'Gráfica', 'Medios vecinales', 'Cine']


def sin_tildes(s):
    return ''.join(c for c in unicodedata.normalize('NFD', s) if unicodedata.category(c) != 'Mn')


def rubro_de(t):
    t = sin_tildes(t.lower())
    if '2017 total' in t: return 'MIX'
    if 'tv abierta' in t: return 'TV abierta'
    if 'tv cable' in t: return 'TV cable'
    if re.search(r'\btv\b', t): return 'TV'          # 2018-2019 (columna tipo); 2022 es cable
    if 'via publica' in t: return 'Vía pública'
    if 'web' in t: return 'Web'
    if 'radio' in t: return 'Radio'
    if 'grafic' in t: return 'Gráfica'
    if 'vecinal' in t: return 'Medios vecinales'
    if 'cine' in t: return 'Cine'
    raise ValueError(t)


def tipo_rubro(v):
    """Columna "tipo" de 2017 y de TV 2018-2019."""
    k = re.sub(r'[^A-Z ]', '', sin_tildes(str(v or '')).upper())
    for clave, r in [('ABIERTA', 'TV abierta'), ('PAGO', 'TV cable'), ('CABLE', 'TV cable'), ('VIA PUB', 'Vía pública'),
                     ('WEB', 'Web'), ('RADIO', 'Radio'), ('GR', 'Gráfica'), ('VECINAL', 'Medios vecinales'), ('CINE', 'Cine')]:
        if clave in k: return r
    raise ValueError('tipo desconocido: ' + repr(v))


def inventario():
    files = {f[:8]: f for f in os.listdir(RAW) if not f.startswith('_')}
    inv = defaultdict(lambda: {'csv': [], 'xlsx': []})
    for l in open(os.path.join(RAW, '_recursos.tsv'), encoding='utf-8'):
        rid, t = l.rstrip('\n').split('\t')
        f = files.get(rid[:8])
        if not f: continue
        y = int(re.search(r'20\d\d', t).group())
        r = rubro_de(t)
        if r == 'TV' and y == 2022: r = 'TV cable'   # el recurso "TV 2022" es el archivo de cable
        ext = 'xlsx' if f.lower().endswith('.xlsx') else 'csv'
        if f not in inv[(y, r)][ext]: inv[(y, r)][ext].append(f)
    return inv


# ---------- lectura ----------
def decode(b):
    for e in ('utf-8-sig', 'cp1252'):
        try: return b.decode(e)
        except UnicodeDecodeError: pass
    return b.decode('latin-1')


def csv_rows(path):
    t = decode(open(path, 'rb').read())
    first = t.splitlines()[0]
    sep = max([';', ',', '\t'], key=first.count)
    rows = []
    for r in csv.reader(io.StringIO(t), delimiter=sep):
        if len(r) == 1 and sep in r[0]:          # filas enteras entre comillas (2022)
            r = next(csv.reader([r[0]], delimiter=sep))
        rows.append(r)
    return rows


def xlsx_rows(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = []
    for r in wb.worksheets[0].iter_rows(values_only=True):
        if r is None or all(v is None or str(v).strip() == '' for v in r):
            if len(out) > 5 and len(out) > 0: pass
            continue
        out.append(list(r))
    return out


def cols(header):
    h = [sin_tildes(str(c or '')).lower().strip() for c in header]
    def find(*keys):
        for i, c in enumerate(h):
            if any(k in c for k in keys): return i
        return None
    return {'fecha': find('fecha'), 'medio': find('medio'), 'importe': find('importe', 'monto'),
            'tipo': find('tipo'), 'razon': find('razon'), 'campana': find('campana')}


def num_fmt(vals):
    """Deduce el formato de los importes de un archivo: 'us' (1,234.56), 'ar' (1.234,56) o 'plain'."""
    us = ar = 0
    for v in vals:
        s = re.sub(r'[^\d.,]', '', v)
        if '.' in s and ',' in s:
            if s.rfind('.') > s.rfind(','): us += 1
            else: ar += 1
        elif re.fullmatch(r'\d{1,3}(,\d{3})+', s): us += 1
        elif re.fullmatch(r'\d{1,3}(\.\d{3})+', s): ar += 1
        elif re.fullmatch(r'\d+,\d{1,2}', s): ar += 1
    return 'us' if us > ar else 'ar' if ar > us else 'plain'


def to_num(v, fmt):
    if isinstance(v, (int, float)): return float(v)
    s = re.sub(r'[^\d.,\-]', '', str(v or ''))
    if not s or not re.search(r'\d', s): return None
    if fmt == 'us': s = s.replace(',', '')
    elif fmt == 'ar': s = s.replace('.', '').replace(',', '.')
    else: s = s.replace(',', '.') if s.count(',') == 1 and '.' not in s else s.replace(',', '')
    try: return float(s)
    except ValueError: return None


def date_order(vals):
    """'dm' o 'md' para fechas con barras, según qué componente supera 12."""
    a = b = 0
    for v in vals:
        m = re.match(r'\s*(\d{1,2})/(\d{1,2})/(\d{2,4})', v)
        if not m: continue
        if int(m.group(1)) > 12: a += 1
        if int(m.group(2)) > 12: b += 1
    return 'dm' if a >= b else 'md'


def to_date(v, order, year):
    if isinstance(v, dt.datetime): return v.date()
    if isinstance(v, dt.date): return v
    if isinstance(v, (int, float)) and 30000 < v < 60000:
        return (dt.datetime(1899, 12, 30) + dt.timedelta(days=int(v))).date()
    s = str(v or '').strip()
    m = re.match(r'(\d{4})-(\d{1,2})-(\d{1,2})', s)
    if m: return dt.date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    m = re.match(r'(\d{1,2})/(\d{1,2})/(\d{2,4})', s)
    if m:
        p, q, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
        if y < 100: y += 2000
        d, mo = (p, q) if order == 'dm' else (q, p)
        try: return dt.date(y, mo, d)
        except ValueError: return dt.date(y, d, mo)
    return None


def leer(path, year, rubro):
    rows = csv_rows(path) if path.lower().endswith('.csv') else xlsx_rows(path)
    hi = next(i for i, r in enumerate(rows) if any('fecha' in sin_tildes(str(c or '')).lower() for c in r))
    c = cols(rows[hi])
    body = [r for r in rows[hi + 1:] if len(r) > max(c['fecha'], c['medio'], c['importe'])]
    fmt = num_fmt([str(r[c['importe']]) for r in body if isinstance(r[c['importe']], str)])
    order = date_order([str(r[c['fecha']]) for r in body if isinstance(r[c['fecha']], str)])
    out = []
    for r in body:
        imp = to_num(r[c['importe']], fmt)
        medio = str(r[c['medio']] or '').strip()
        if imp is None or not medio: continue
        f = to_date(r[c['fecha']], order, year)
        ru = tipo_rubro(r[c['tipo']]) if rubro in ('MIX', 'TV') else rubro
        o = {'anio': year, 'fecha': f, 'rubro': ru, 'medio_fuente': medio, 'importe': imp}
        if c['campana'] is not None: o['campana'] = str(r[c['campana']] or '').strip()
        out.append(o)
    return out


# ---------- medios ----------
def limpiar(s):
    """Saca los caracteres de control que deja la doble codificación del archivo 2017 ("CLARÍ\x8dN")."""
    return re.sub(r'\s+', ' ', re.sub(r'[\x80-\x9f]', '', s or '')).strip()


RADIO_RELLENO = {'AM', 'FM', 'RADIO', 'LA', 'EL', 'DE', 'DEL', 'Y', '-'}
ALIAS = {('Gráfica', 'CRONICA MATUTINA'): 'CRONICA', ('Gráfica', 'AMBITO'): 'AMBITO FINANCIERO'}


def clave(medio, rubro):
    s = sin_tildes(limpiar(medio)).upper()
    if rubro == 'Web':
        s = re.sub(r'^(HTTPS?://)?(WWW\.)?', '', s).rstrip('/ ')
    s = re.sub(r'[\s.,;:\-]+$', '', s)
    s = re.sub(r'\s+(S\.?A\.?|S\.?R\.?L\.?)$', '', s)
    if rubro == 'Radio':
        # "RADIO MITRE AM 790" (2017-2019) y "AM 790 RADIO MITRE" (2020+) son la misma emisora
        m = re.search(r'\b(AM|FM)\s*(\d{2,4}(?:[.,]\d)?)\b', s) or re.search(r'\b(\d{2,4}(?:[.,]\d))\s*(FM)\b', s)
        if m:
            band, freq = (m.group(1), m.group(2)) if m.group(1) in ('AM', 'FM') else (m.group(2), m.group(1))
            resto = re.sub(r'[^A-Z0-9 ]', ' ', (s[:m.start()] + ' ' + s[m.end():]))
            toks = sorted(t for t in resto.split() if t not in RADIO_RELLENO and t != freq)
            s = f"{band} {freq.replace(',', '.')} {' '.join(toks)}".strip()
    if rubro == 'Gráfica' and 'SUPLEMENTO' in s and 'AGENDA' in s:
        return 'SUPLEMENTO AGENDA CULTURAL'  # "Disfrutemos BA" hasta 2023, "Descubrir BA" desde 2024
    if rubro == 'Gráfica':
        s = re.sub(r'^(DIARIO|REVISTA|PERIODICO|SEMANARIO)\s+', '', s)
    return ALIAS.get((rubro, s), s)


# Grupos de medios (asignación orientativa, solo los casos conocidos): (grupo, rubros o None, regex sobre la clave)
GRUPOS = [
    ('Grupo Clarín', None, r'CLARIN|^OLE$|OLE\.COM|^CANAL 13$|^TN$|^TN - CANAL 13$|TN\.COM|ARTEAR|ELTRECE|^AM 790 MITRE$|^FM 99\.9 100$'),
    ('La Nación', None, r'LA NACION|^LN\+|LANACION\.COM|^OHLALA|^LUGARES$|^LIVING$|^JARDIN$'),
    ('Grupo Perfil', None, r'^PERFIL$|PERFIL\.COM|^NOTICIAS$|^CARAS$|^FORTUNA$|^WEEKEND$|^NET TV$'),
    ('Telefe', None, r'TELEFE'),
    ('Grupo América', None, r'^AMERICA( TV)?$|^AMERICA 24$|^A24$|^AMERICA - A24$|^AMERICA 2$|^AM 910 RED$'),
    ('Grupo Indalo', None, r'^C5N$|^AM 710 DIEZ$|^FM 101\.5 POP$|^FM 98\.3 MEGA$|MINUTOUNO'),
    ('Grupo Crónica', None, r'^CRONICA( TV)?$|CRONICA\.COM'),
    ('Grupo Octubre', None, r'PAGINA ?12|^AM 750\b'),
    ('Google', ('Web',), r'GOOGLE|YOUTUBE|DV360|^WAZE'),
    ('Meta', ('Web',), r'FACEBOOK|INSTAGRAM|^META$|WHATSAPP'),
    ('Disney', ('TV cable', 'Web'), r'^ESPN|STAR CHANNEL|FOX SPORTS|DISNEY'),
    ('Warner Bros. Discovery', ('TV cable', 'Web'), r'DISCOVERY|^TNT|^CNN$|^HBO|WARNER'),
]


def grupo_de(rubro, k):
    for i, (g, rubros, rx) in enumerate(GRUPOS):
        if (rubros is None or rubro in rubros) and re.search(rx, k):
            return i
    return -1


def ipc_mensual():
    wb = openpyxl.load_workbook(IPC_XLSX, read_only=True, data_only=True)
    ipc = {}
    for r in wb.worksheets[0].iter_rows(values_only=True):
        if isinstance(r[0], dt.datetime) and isinstance(r[1], (int, float)):
            ipc[(r[0].year, r[0].month)] = float(r[1])
    return ipc


def presupuesto(ipc_base, base):
    """Partida Publicidad y propaganda (data/presupuesto_publicidad.json, de build_presupuesto.py), con valores reales."""
    p = os.path.join(HERE, 'presupuesto_publicidad.json')
    if not os.path.exists(p): return None
    P = json.load(open(p, encoding='utf-8'))
    ipc = {int(k): v for k, v in P['ipc_prom'].items()}
    f = lambda y: ipc[base] / ipc[y]
    return {'fuente': P['fuente'], 'anios': {y: {**a, 'fr': round(f(int(y)), 6)} for y, a in P['anios'].items()}}


def main():
    inv = inventario()
    regs, control, fuentes = [], [], {}
    for (y, r), fs in sorted(inv.items()):
        tot = {}
        for f in fs['csv'] + fs['xlsx']:
            try:
                rr = leer(os.path.join(RAW, f), y, r)
            except Exception as e:
                print(f'  ! {f}: {e}'); continue
            tot[f] = (sum(x['importe'] for x in rr), len(rr), rr)
        con = {f: v for f, v in tot.items() if v[1] > 0}
        if not con:
            print(f'  ! {y} {r}: sin datos'); continue
        # CSV primero; si hay varios iguales (duplicados en BA Data) se usa uno solo
        f = next((f for f in fs['csv'] if f in con), None) or next(f for f in fs['xlsx'] if f in con)
        if len(con) > 1:
            ref = [v[0] for k, v in con.items() if k != f]
            d = max(abs(con[f][0] / x - 1) for x in ref) if all(ref) else 0
            control.append((y, r, f, con[f][0], d))
            if d > 0.001: print(f'  ! {y} {r}: CSV y XLSX difieren {d:.2%}  ({[ (k, round(v[0])) for k, v in con.items()]})')
        regs += con[f][2]
        fuentes[f'{y}|{r}'] = f
    print(f'{len(regs):,} órdenes leídas')

    ipc = ipc_mensual()
    ult = max(y for y in range(2012, 2100) if all((y, m) in ipc for m in range(1, 13)))
    years = sorted({x['anio'] for x in regs})
    base = min(ult, max(years))
    ipc_base = sum(ipc[(base, m)] for m in range(1, 13)) / 12

    def ipc_de(x):
        f = x['fecha']
        if f and f.year == x['anio'] and (f.year, f.month) in ipc: return ipc[(f.year, f.month)]
        ms = [ipc[(x['anio'], m)] for m in range(1, 13) if (x['anio'], m) in ipc]
        return sum(ms) / len(ms)

    # medios: clave normalizada + nombre más frecuente
    nombres = defaultdict(Counter)
    for x in regs:
        x['medio'] = clave(x['medio_fuente'], x['rubro'])
        nombres[(x['rubro'], x['medio'])][limpiar(x['medio_fuente'])] += x['anio'] - 2000  # gana la grafía reciente
        x['real'] = x['importe'] * ipc_base / ipc_de(x)
        if x['fecha'] and x['fecha'].year != x['anio']:
            x['fecha_ok'] = False

    with gzip.open(os.path.join(HERE, 'pauta.csv.gz'), 'wt', encoding='utf-8', newline='') as g:
        w = csv.writer(g)
        w.writerow(['anio', 'fecha', 'rubro', 'medio', 'medio_fuente', 'importe', 'importe_real_' + str(base)])
        for x in regs:
            w.writerow([x['anio'], x['fecha'] or '', x['rubro'], x['medio'], x['medio_fuente'], round(x['importe'], 2), round(x['real'], 2)])

    # ---------- agregados ----------
    publ = defaultdict(set)
    for x in regs: publ[x['anio']].add(x['rubro'])
    medios = sorted({(x['rubro'], x['medio']) for x in regs})
    midx = {m: i for i, m in enumerate(medios)}
    agg = defaultdict(lambda: [0.0, 0.0, 0])  # (anio, mes, medio) -> nominal, real, órdenes
    for x in regs:
        mes = x['fecha'].month if x['fecha'] and x['fecha'].year == x['anio'] else 0
        a = agg[(x['anio'], mes, midx[(x['rubro'], x['medio'])])]
        a[0] += x['importe']; a[1] += x['real']; a[2] += 1
    data = {
        'generado': dt.datetime.now().isoformat(timespec='seconds'),
        'fuente': 'BA Data · data.buenosaires.gob.ar/dataset/pauta-publicitaria',
        'years': years, 'base': base, 'rubros': RUBROS,
        'publicado': {str(y): sorted(publ[y], key=RUBROS.index) for y in years},
        'ipc_anual': {str(y): round(sum(ipc[(y, m)] for m in range(1, 13) if (y, m) in ipc) / max(1, sum((y, m) in ipc for m in range(1, 13))), 4) for y in years},
        'grupos': [g[0] for g in GRUPOS],
        # medios: [rubro, nombre, grupo (-1 = sin asignar)]
        'medios': [[RUBROS.index(r), nombres[(r, m)].most_common(1)[0][0], grupo_de(r, m)] for r, m in medios],
        'presupuesto': presupuesto(ipc_base, base),
        # filas: [año, mes (0 = sin fecha), medio, nominal, real, órdenes]
        'rows': [[y, m, i, round(v[0]), round(v[1]), v[2]] for (y, m, i), v in sorted(agg.items())],
    }
    with open(OUT_JSON, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, separators=(',', ':'))
    with open(OUT_JSON[:-5] + '.js', 'w', encoding='utf-8') as f:  # para abrir index.html sin servidor
        f.write('window.PAUTA_DATA=' + json.dumps(data, ensure_ascii=False, separators=(',', ':')) + ';\n')
    print(f'{OUT_JSON}: {os.path.getsize(OUT_JSON) / 1e6:.2f} MB | {len(medios):,} medios | base real {base}')

    # resumen
    tab = defaultdict(float)
    for x in regs: tab[(x['anio'], x['rubro'])] += x['importe']
    print('\nNominal, millones de $')
    print('año   ' + ''.join(f'{r[:9]:>11s}' for r in RUBROS) + '      TOTAL     real')
    for y in years:
        tr = sum(x['real'] for x in regs if x['anio'] == y)
        print(f'{y}  ' + ''.join(f'{tab[(y, r)] / 1e6:11,.0f}' if (y, r) in tab else f'{"—":>11s}' for r in RUBROS) + f'{sum(tab[(y, r)] for r in RUBROS) / 1e6:11,.0f}{tr / 1e6:9,.0f}')
    malas = sum(1 for x in regs if x.get('fecha_ok') is False)
    sinf = sum(1 for x in regs if not x['fecha'])
    print(f'\nfechas fuera del año: {malas} | sin fecha: {sinf}')
    for y, r, f, t, d in control:
        print(f'  control {y} {r:17s} {d:.4%}  {f}')


if __name__ == '__main__':
    main()
